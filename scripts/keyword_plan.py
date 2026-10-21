#!/usr/bin/env python3
"""Write docs/keyword-plan.md from keywords/Dallas-Power-Washing-Keywords.xlsx.

Every keyword with 10+ searches a month must be ASSIGNED (target or supporting keyword of a page)
or UNUSED with a one-line reason. The script fails if any 10+ keyword is missing from the map.
Keywords under 10 searches are summarized by group, not listed one by one.
"""
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "keywords" / "Dallas-Power-Washing-Keywords.xlsx"

T, S = "target", "supporting"
ASSIGN = {
    # home
    "power washing dallas": ("/", T),
    "pressure washing in dallas": ("/", T),
    "pressure washing dallas": ("/", T),
    "pressure washing dallas texas": ("/", T),
    "pressure washing in dallas tx": ("/", S),
    "pressure washing dfw": ("/", S),
    "power washing services near me": ("/", S),
    "power washing in my area": ("/", S),
    "power washing service in my area": ("/", S),
    "who does pressure washing in my area": ("/how-it-works/", S),
    "pressure washing companies dallas texas": ("/", S),
    "pressure washing places": ("/service-areas/", S),
    "pressure washing needed near me": ("/pressure-washing/", S),
    # services
    "pressure washing description": ("/pressure-washing/", S),
    "pressure washing facts": ("/pressure-washing/", S),
    "dangers of pressure washing": ("/pressure-washing/", S),
    "what to use before pressure washing": ("/pressure-washing/", S),
    "house washing services near me": ("/house-washing/", S),
    "house washing dallas": ("/house-washing/", T),
    "wash house dallas": ("/house-washing/", S),
    "soft pressure washing near me": ("/house-washing/", S),
    "what is soft wash pressure washing": ("/house-washing/", S),
    "soft wash vs pressure washing": ("/guides/power-washing-vs-pressure-washing/", S),
    "roof cleaning dallas": ("/roof-soft-wash/", T),
    "roof cleaning how much": ("/roof-soft-wash/", S),
    "how often roof cleaning": ("/roof-soft-wash/", S),
    "driveway cleaning dallas": ("/driveway-cleaning/", T),
    "driveway dallas": ("/driveway-cleaning/", S),
    "how often pressure wash driveway": ("/driveway-cleaning/", S),
    "commercial pressure washing dallas": ("/power-washing/", T),
    "commercial pressure washing companies near me": ("/power-washing/", S),
    # guides
    "how much does pressure washing cost": ("/guides/pressure-washing-cost/", T),
    "how much does pressure washing cost per square foot": ("/guides/pressure-washing-cost/", S),
    "how much is pressure washing per square foot": ("/guides/pressure-washing-cost/", S),
    "going rate for pressure washing": ("/guides/pressure-washing-cost/", S),
    "what is a good price for pressure washing": ("/guides/pressure-washing-cost/", S),
    "how much is a pressure washing job": ("/guides/pressure-washing-cost/", S),
    "house washing cost": ("/guides/pressure-washing-cost/", S),
    "how much is house washing": ("/guides/pressure-washing-cost/", S),
    "how often pressure wash house": ("/guides/how-often-pressure-wash-house/", T),
    "how often to power wash a house": ("/guides/how-often-pressure-wash-house/", S),
    "how often to pressure wash your house": ("/guides/how-often-pressure-wash-house/", S),
    "what is the difference between power washing and pressure washing": ("/guides/power-washing-vs-pressure-washing/", T),
    "commercial pressure washing rates": ("/guides/commercial-pressure-washing-cost/", T),
    "how much does commercial pressure washing cost": ("/guides/commercial-pressure-washing-cost/", S),
    "commercial pressure washing price list": ("/guides/commercial-pressure-washing-cost/", S),
    "do you need a license to pressure wash in texas": ("/guides/pressure-washing-license-texas/", T),
    # towns
    "pressure washing fort worth": ("/fort-worth-pressure-washing/", T),
    "pressure washing frisco": ("/frisco-pressure-washing/", T),
    "pressure washing denton": ("/denton-pressure-washing/", T),
    "denton pressure washing": ("/denton-pressure-washing/", S),
    "pressure washing plano": ("/plano-pressure-washing/", T),
    "power washing plano": ("/plano-pressure-washing/", S),
    "pressure washing arlington tx": ("/arlington-pressure-washing/", T),
    "pressure washing arlington": ("/arlington-pressure-washing/", S),
    "mckinney pressure washing": ("/mckinney-pressure-washing/", T),
    "pressure washing mckinney": ("/mckinney-pressure-washing/", S),
    "pressure washing allen tx": ("/allen-pressure-washing/", T),
    "allen pressure washing": ("/allen-pressure-washing/", S),
}

UNUSED = {
    "pressure wash car wash near me": "Car wash search (vehicles), not exterior cleaning of property.",
    "chemicals for pressure washing": "DIY chemical shopping search; the site gives no chemical instructions.",
    "tools for pressure washing": "DIY equipment search.",
    "house cleaning jobs dallas tx": "Job search for maid work.",
    "how much to charge for power washing": "Business-owner pricing search, not a customer.",
    "power washing salary": "Job and salary search.",
    "pressure washing instructions": "DIY instruction search; the site gives no machine instructions.",
    "house cleaning dallas texas": "Maid and house cleaning (indoor), a service the line does not handle.",
    "wash house hours": "Laundromat search.",
    "driveway cleaning houston": "Houston, outside DFW.",
    "house washing houston": "Houston, outside DFW.",
    "commercial pressure washing equipment near me": "Equipment shopping search.",
    "pressure washing plus": "Names a company.",
    "pressure washing friendswood tx": "Houston area, outside DFW.",
    "pressure washing deer park": "Houston area, outside DFW.",
    "driveway repair dallas tx": "Concrete repair, a service the line does not handle.",
    "washing machine dallas": "Appliance search.",
    "how to get commercial pressure washing jobs": "Business-owner search.",
    "pressure washing pflugerville": "Austin area, outside DFW.",
    "what psi needed to pressure wash house": "PSI setting search; the site never gives machine settings.",
    "how much pressure to pressure wash a house": "PSI setting search; the site never gives machine settings.",
    "how much to charge for a pressure washing job": "Business-owner pricing search.",
    "how many driveways can a house have": "Zoning question unrelated to cleaning.",
    "what are the best commercial pressure washers": "Equipment shopping search.",
    "are commercial washing machines better": "Laundry appliance search.",
    "arlington car wash near me": "Car wash search.",
    "pressure washer rental dallas": "Equipment rental, a service the line does not handle.",
    "pressure washing dallas georgia": "Dallas, Georgia, another state.",
    "dallas house cleaning prices": "Maid service prices (indoor cleaning).",
    "house washing katy": "Houston area, outside DFW.",
    "house washing york pa": "Another state.",
    "roof cleaning league city": "Houston area, outside DFW.",
    "roof cleaning el paso": "El Paso, outside DFW.",
    "roof cleaning katy": "Houston area, outside DFW.",
    "pressure washing arlington va": "Arlington, Virginia, another state.",
    "pressure washing tricks": "DIY technique search.",
    "pressure washing fulshear": "Houston area, outside DFW.",
    "pressure washing flowery branch ga": "Another state.",
    "pressure washing mount vernon": "Not in DFW (other states).",
}


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    all_rows = {}
    for r in list(wb["ALL"].iter_rows(values_only=True))[1:]:
        if r[0]:
            all_rows[r[0]] = {"vol": r[1] or 0, "sd": r[2], "cpc": r[3], "seeds": r[4]}
    types = defaultdict(set)
    for ws in wb.worksheets:
        if ws.title == "ALL":
            continue
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if r[0]:
                types[r[0]].add(r[1])
    missing = [k for k, v in all_rows.items() if v["vol"] >= 10 and k not in ASSIGN and k not in UNUSED]
    if missing:
        print("UNMAPPED 10+ keywords:", missing)
        return 1
    stray = [k for k in list(ASSIGN) + list(UNUSED) if k not in all_rows]
    if stray:
        print("Mapped keywords not in the file:", stray)
        return 1
    by_page = defaultdict(list)
    for k, (page, role) in ASSIGN.items():
        by_page[page].append((role != T, -all_rows[k]["vol"], k, role))
    ten_plus = sum(1 for v in all_rows.values() if v["vol"] >= 10)
    out = [
        "# Keyword plan: dallaspowerwashingpro.com",
        "",
        "Source: keywords/Dallas-Power-Washing-Keywords.xlsx (Ubersuggest, United States, pulled Oct 21 2026).",
        f"{len(all_rows)} unique keywords. {ten_plus} have 10 or more searches a month. Every one of those {ten_plus} is",
        "either assigned to a page below (as its target or a supporting keyword) or listed under Unused with a reason.",
        "This file is generated by `python3 scripts/keyword_plan.py`, which fails if any 10+ keyword is left out.",
        "",
        "Notes:",
        "- \"Near me\", \"in my area\" and national question terms are national volume, not Dallas volume. They are",
        "  used only as supporting keywords or as guide targets, never as a page's main local target.",
        "- The brand says POWER washing, so the home page title and H1 lead with power washing and also cover",
        "  pressure washing (both 720/mo).",
        "- Volume is monthly US searches; SD is Ubersuggest SEO difficulty (0 to 100).",
        "",
        "## Assigned keywords by page",
        "",
    ]
    for page in sorted(by_page, key=lambda p: (p != "/", p)):
        out += [f"### {page}", "", "| Keyword | Role | Volume | SD | Type |", "|---|---|---:|---:|---|"]
        for _, _, k, role in sorted(by_page[page]):
            v = all_rows[k]
            out.append(f"| {k} | {role} | {v['vol']} | {v['sd']} | {', '.join(sorted(t for t in types[k] if t)) or '-'} |")
        out.append("")
    out += ["## Pages with no 10+ keyword in this pull", "",
            "- /fence-cleaning/ and /deck-cleaning/: existing indexed service pages. The pull has only zero-volume",
            "  rows for them (\"how often pressure wash fence\", \"how often to pressure wash a deck\"). Kept as real",
            "  service pages because callers ask for these surfaces; listed for fresh data.",
            "- /addison-pressure-washing/, /carrollton-pressure-washing/, /irving-pressure-washing/: live town pages",
            "  with no keyword in this pull. Kept at their URLs as shorter, honest town pages.",
            "- /flower-mound-pressure-washing/: the seed \"pressure washing flower mound\" returned 0 searches. Kept as",
            "  a shorter town page at its URL.",
            "- /house-washing/ is a NEW page: \"house washing services near me\" 1,300 and \"soft pressure washing near",
            "  me\" 170 justify a dedicated house and soft washing page.",
            "",
            "## Unused keywords (10+ searches) and why", "", "| Keyword | Volume | Reason |", "|---|---:|---|"]
    for k in sorted(UNUSED, key=lambda k: -all_rows[k]["vol"]):
        out.append(f"| {k} | {all_rows[k]['vol']} | {UNUSED[k]} |")
    zero = [k for k, v in all_rows.items() if v["vol"] < 10]
    out += ["", f"## Keywords under 10 searches ({len(zero)})", "",
            "Summarized by group, not assigned one by one: zero-volume Dallas, Fort Worth, Arlington, Frisco, Plano, Allen,",
            "McKinney, Denton and Flower Mound phrasings (covered naturally by the matching page), company-name rows",
            "(not used), \"how often\" variants for houses, driveways, decks, fences and roofs (covered by the frequency",
            "guide and service pages), ZIP code rows from the Houston area and other cities, household-count and",
            "house-warming rows (off topic), and DIY rows (not used).",
            "",
            "## Pages that would gain most from fresh Ubersuggest data",
            "",
            "1. /fence-cleaning/ and /deck-cleaning/ (no measurable volume in this pull).",
            "2. Addison, Carrollton, Irving and Flower Mound town pages (pull a seed per town).",
            "3. Possible new town pages: Richardson, Garland, Grand Prairie, Mesquite, Lewisville, Southlake, Keller",
            "   and Prosper were not seeded. Any with 10+ searches a month would get a page under Maurice's rule.",
            "4. /power-washing/ (commercial): seed \"parking lot cleaning dallas\" and \"storefront pressure washing\".",
            ""]
    (ROOT / "docs" / "keyword-plan.md").write_text("\n".join(out))
    print(f"wrote docs/keyword-plan.md: {len(ASSIGN)} assigned, {len(UNUSED)} unused, {ten_plus} with 10+")
    return 0


if __name__ == "__main__":
    sys.exit(main())
